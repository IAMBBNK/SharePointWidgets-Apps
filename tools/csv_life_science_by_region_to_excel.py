#!/usr/bin/env python3
"""
Split a CSV into region Excel workbooks for Life Science entries.

Filters Sector = Life Science, skips blank Site Codes, then for each region
(APAC / NASA / EMEA) writes workbooks with up to 15 site sheets each.

Each site sheet is a vertical Field / Value table (transposed).

Files are named:
  LifeScience_NASA_1.xlsx
  LifeScience_NASA_2.xlsx
  ...

Examples:
  python tools/csv_life_science_by_region_to_excel.py "tools/Site Master Data (2).csv"
  python tools/csv_life_science_by_region_to_excel.py data/export.csv -o output/ls --max-sheets 15
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

import pandas as pd
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

REGIONS = ("APAC", "NASA", "EMEA")
DEFAULT_SECTOR_VALUE = "Life Science"
DEFAULT_MAX_SHEETS = 15
SHEET_NAME_MAX = 31
INVALID_SHEET_CHARS = re.compile(r"[\\/*?:\[\]]+")


def find_column(columns: list[str], candidates: list[str]) -> str | None:
    """Return the first matching column name (case-insensitive, stripped)."""
    normalized = {str(c).strip().lower(): c for c in columns}
    for cand in candidates:
        hit = normalized.get(cand.strip().lower())
        if hit is not None:
            return hit
    return None


def sanitize_sheet_name(name: str, used: set[str]) -> str:
    """Excel sheet names: max 31 chars, no \\ / * ? : [ ]."""
    cleaned = INVALID_SHEET_CHARS.sub("-", str(name or "").strip())
    cleaned = re.sub(r"\s+", " ", cleaned).strip(" .'\"") or "Entry"
    cleaned = cleaned[:SHEET_NAME_MAX]

    base = cleaned
    n = 2
    while cleaned.lower() in {u.lower() for u in used}:
        suffix = f" ({n})"
        cleaned = (base[: SHEET_NAME_MAX - len(suffix)] + suffix)[:SHEET_NAME_MAX]
        n += 1

    used.add(cleaned)
    return cleaned


def pick_sheet_title(row: pd.Series, columns: list[str]) -> str:
    """Prefer Site Code / Site Name / Title for the sheet tab name."""
    preferred = [
        "Site Code",
        "Site ID",
        "SiteId",
        "Site Name",
        "SiteName",
        "Title",
        "Name",
    ]
    for col_name in preferred:
        col = find_column(columns, [col_name])
        if col is None:
            continue
        val = row.get(col)
        if pd.notna(val) and str(val).strip():
            return str(val).strip()

    for col in columns:
        val = row.get(col)
        if pd.notna(val) and str(val).strip():
            return str(val).strip()
    return "Entry"


def chunk_dataframe(df: pd.DataFrame, size: int) -> list[pd.DataFrame]:
    if df.empty:
        return []
    return [df.iloc[i : i + size].copy() for i in range(0, len(df), size)]


def row_to_vertical_frame(row: pd.Series) -> pd.DataFrame:
    """Transpose one site row into a vertical Field / Value table."""
    fields: list[str] = []
    values: list[str] = []
    for col, val in row.items():
        fields.append(str(col))
        if pd.isna(val):
            values.append("")
        else:
            values.append(str(val))
    return pd.DataFrame({"Field": fields, "Value": values})


def autosize_columns(ws, df: pd.DataFrame) -> None:
    for idx, col in enumerate(df.columns, start=1):
        series = df.iloc[:, idx - 1].astype(str)
        max_len = max([len(str(col))] + [len(v) for v in series.tolist()])
        ws.column_dimensions[get_column_letter(idx)].width = min(max_len + 2, 60)


def write_workbook(
    df: pd.DataFrame,
    region: str,
    part: int,
    output_path: Path,
) -> int:
    """Write one vertical sheet per row. Returns number of sheets created."""
    if df.empty:
        return 0

    used_names: set[str] = set()
    columns = list(df.columns)

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for i, (_, row) in enumerate(df.iterrows(), start=1):
            title = pick_sheet_title(row, columns)
            sheet_name = sanitize_sheet_name(title, used_names)
            frame = row_to_vertical_frame(row)
            frame.to_excel(writer, sheet_name=sheet_name, index=False)

            ws = writer.book[sheet_name]
            autosize_columns(ws, frame)

            end_col = get_column_letter(frame.shape[1])
            ref = f"A1:{end_col}{frame.shape[0] + 1}"
            table_name = f"T_{region}_{part}_{i}"
            table_name = re.sub(r"[^A-Za-z0-9_]", "_", table_name)[:250]
            table = Table(displayName=table_name, ref=ref)
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
            ws.add_table(table)

    return len(df)


def write_region_books(
    region_df: pd.DataFrame,
    region: str,
    site_code_col: str,
    out_dir: Path,
    max_sheets: int,
) -> list[tuple[str, int]]:
    """Split a region into numbered workbooks of max_sheets sheets each."""
    results: list[tuple[str, int]] = []
    if region_df.empty:
        return results

    work = region_df.copy()
    work["_site_code"] = work[site_code_col].astype(str).str.strip()
    work = work.sort_values(by=["_site_code"], kind="mergesort")
    clean = work.drop(columns=["_site_code"])

    chunks = chunk_dataframe(clean, max_sheets)
    for part_idx, chunk in enumerate(chunks, start=1):
        name = f"LifeScience_{region}_{part_idx}.xlsx"
        path = out_dir / name
        count = write_workbook(chunk, region, part_idx, path)
        results.append((name, count))

    return results


def main() -> int:
    parser = argparse.ArgumentParser(
        description=(
            "Filter CSV for Sector=Life Science and write Excel workbooks "
            "per region (~15 vertical site sheets each)."
        )
    )
    parser.add_argument("csv_path", type=Path, help="Path to the input CSV")
    parser.add_argument(
        "-o",
        "--output-dir",
        type=Path,
        default=None,
        help="Output folder (default: <csv_stem>_life_science_excel next to the CSV)",
    )
    parser.add_argument(
        "--max-sheets",
        type=int,
        default=DEFAULT_MAX_SHEETS,
        help=f"Max site sheets per workbook (default: {DEFAULT_MAX_SHEETS})",
    )
    parser.add_argument("--sector-col", default=None)
    parser.add_argument("--region-col", default=None)
    parser.add_argument(
        "--sector-value",
        default=DEFAULT_SECTOR_VALUE,
        help=f'Sector value to keep (default: "{DEFAULT_SECTOR_VALUE}")',
    )
    parser.add_argument("--encoding", default="utf-8-sig")
    parser.add_argument("--delimiter", default=None)
    args = parser.parse_args()

    if args.max_sheets < 1:
        print("--max-sheets must be >= 1", file=sys.stderr)
        return 1

    csv_path: Path = args.csv_path
    if not csv_path.is_file():
        print(f"CSV not found: {csv_path}", file=sys.stderr)
        return 1

    read_kwargs = {"dtype": str, "encoding": args.encoding, "keep_default_na": False}
    if args.delimiter:
        read_kwargs["sep"] = args.delimiter

    try:
        df = pd.read_csv(csv_path, **read_kwargs)
    except Exception as exc:  # noqa: BLE001
        print(f"Failed to read CSV: {exc}", file=sys.stderr)
        return 1

    if df.empty:
        print("CSV has no rows.", file=sys.stderr)
        return 1

    columns = [str(c) for c in df.columns]

    sector_col = args.sector_col or find_column(
        columns, ["Sector", "Business Sector", "sector"]
    )
    region_col = args.region_col or find_column(
        columns, ["Region", "Geo", "Geography", "region"]
    )
    site_code_col = find_column(
        columns, ["Site Code", "SiteCode", "Site ID", "SiteId"]
    )

    if not sector_col:
        print(
            "Could not find a Sector column. Pass --sector-col. "
            f"Available: {', '.join(columns)}",
            file=sys.stderr,
        )
        return 1
    if not region_col:
        print(
            "Could not find a Region column. Pass --region-col. "
            f"Available: {', '.join(columns)}",
            file=sys.stderr,
        )
        return 1
    if not site_code_col:
        print(
            "Could not find a Site Code column (required). "
            f"Available: {', '.join(columns)}",
            file=sys.stderr,
        )
        return 1

    sector_norm = df[sector_col].astype(str).str.strip().str.casefold()
    target_sector = args.sector_value.strip().casefold()
    ls = df[sector_norm == target_sector].copy()
    ls_before_site_filter = len(ls)

    site_ok = ls[site_code_col].astype(str).str.strip()
    ls = ls[site_ok.ne("") & site_ok.str.casefold().ne("nan")].copy()
    skipped_no_code = ls_before_site_filter - len(ls)

    ls["_region_key"] = ls[region_col].astype(str).str.strip().str.casefold()

    out_dir = args.output_dir
    if out_dir is None:
        out_dir = csv_path.with_name(f"{csv_path.stem}_life_science_excel")
    out_dir.mkdir(parents=True, exist_ok=True)

    for old in out_dir.glob("LifeScience_*.xlsx"):
        old.unlink()

    print(f"Input: {csv_path}")
    print(f"Sector column: {sector_col!r} = {args.sector_value!r}")
    print(f"Region column: {region_col!r}")
    print(f"Site Code column: {site_code_col!r}")
    print(f"Max sheets per book: {args.max_sheets}")
    print(f"Layout: vertical Field / Value")
    print(f"Life Science rows (before site-code filter): {ls_before_site_filter}")
    if skipped_no_code:
        print(f"Skipped (no Site Code): {skipped_no_code}")
    print(f"Life Science rows kept: {len(ls)}")
    print(f"Output folder: {out_dir}")

    for region in REGIONS:
        region_df = ls[ls["_region_key"] == region.casefold()].drop(
            columns=["_region_key"]
        )
        print(f"\n{region} ({len(region_df)} sites):")
        written = write_region_books(
            region_df, region, site_code_col, out_dir, args.max_sheets
        )
        if not written:
            print("  (no files)")
            continue
        for name, count in written:
            print(f"  {count} sheet(s) -> {name}")

    unmatched = ls[~ls["_region_key"].isin({r.casefold() for r in REGIONS})]
    if not unmatched.empty:
        samples = sorted(
            unmatched[region_col].astype(str).str.strip().unique().tolist()
        )
        print(
            f"\nNote: {len(unmatched)} Life Science row(s) skipped "
            f"(region not in APAC/NASA/EMEA). Values: {samples}"
        )

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
